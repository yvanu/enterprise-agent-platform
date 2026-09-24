from io import BytesIO

import pytest
from docx import Document

from app.agents.knowledge.parser import extract_text


def test_extract_text_file():
    assert extract_text("demo.md", "标题\n正文".encode()) == "标题\n正文"


def test_extract_docx():
    buffer = BytesIO()
    document = Document()
    document.add_paragraph("企业知识文档")
    document.add_paragraph("验收要求")
    document.save(buffer)

    text = extract_text("demo.docx", buffer.getvalue())

    assert "企业知识文档" in text
    assert "验收要求" in text


def test_reject_unsupported_file():
    with pytest.raises(ValueError):
        extract_text("demo.exe", b"test")
